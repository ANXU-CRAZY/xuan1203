from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = r"D:\xuan1203\黄河生态方舟系统软著补正修改策略.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(10.5)
    if color:
        r.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, bold=True, color=(255, 255, 255))
        set_cell_shading(hdr[i], "2F5597")
        if widths:
            hdr[i].width = widths[i]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value))
            if widths:
                cells[i].width = widths[i]
    doc.add_paragraph()
    return table


def set_doc_defaults(doc):
    sec = doc.sections[0]
    sec.top_margin = Cm(2.3)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)

    styles = doc.styles
    styles["Normal"].font.name = "宋体"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    styles["Normal"].font.size = Pt(10.5)

    for name, size, color in [
        ("Title", 24, "1F4E79"),
        ("Heading 1", 16, "1F4E79"),
        ("Heading 2", 13, "2F5597"),
        ("Heading 3", 11, "365F91"),
    ]:
        style = styles[name]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)


def add_body_paragraph(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(6)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        r1.font.name = "宋体"
        r1._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        r1.font.size = Pt(10.5)
        rest = text[len(bold_prefix):]
        r2 = p.add_run(rest)
    else:
        r2 = p.add_run(text)
    r2.font.name = "宋体"
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r2.font.size = Pt(10.5)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        r.font.name = "宋体"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        r.font.size = Pt(10.5)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        r.font.name = "宋体"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        r.font.size = Pt(10.5)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(10)
    r.bold = True


def main():
    doc = Document()
    set_doc_defaults(doc)

    # Cover
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("黄河生态方舟系统 V1.0")
    r.font.name = "微软雅黑"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.size = Pt(24)
    r.bold = True
    r.font.color.rgb = RGBColor(31, 78, 121)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("软件著作权登记补正修改策略")
    r.font.name = "微软雅黑"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.size = Pt(22)
    r.bold = True
    r.font.color.rgb = RGBColor(31, 78, 121)

    doc.add_paragraph()
    add_table(
        doc,
        ["项目", "内容"],
        [
            ["补正对象", "《黄河生态方舟系统操作手册》"],
            ["补正原因", "文档部分截图不完整；需提交详细使用说明文档，展现软件使用流程和全部功能；内容需连贯，截图需清晰完整。"],
            ["修改目标", "将原文档由“功能介绍型材料”修改为“可审查、可跟随、流程完整的软件操作手册”。"],
            ["建议交付文件", "黄河生态方舟系统V1.0操作手册_补正版.docx；如平台允许，同时导出PDF版本。"],
        ],
        widths=[Cm(4), Cm(12)],
    )
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("编制日期：2026年5月")
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(11)
    doc.add_page_break()

    doc.add_heading("一、总体修改原则", level=1)
    add_body_paragraph(
        doc,
        "本次补正应围绕审查意见逐项回应，重点不是增加背景介绍，而是补齐“软件如何使用”的实际操作说明。修改后的手册应当让审查人员无需运行系统，也能通过文字和截图理解软件从登录、数据采集、查询展示到后台审核管理的完整流程。"
    )
    add_table(
        doc,
        ["要求", "具体修改方向"],
        [
            ["详细使用说明", "每个功能写明入口、操作步骤、填写字段、保存或查询后的结果。"],
            ["展现使用流程", "按真实使用顺序组织：登录系统 → 进入模块 → 操作数据 → 保存/查询 → 查看结果。"],
            ["展现全部功能", "用户端和管理员端功能都要覆盖，尤其是数据上传、地图展示、后台审核、SHP上传、样线绘制、用户权限和积分商城。"],
            ["内容连贯", "章节之间按角色和业务流程排列，避免零散截图堆叠。"],
            ["截图清晰完整", "删除空白图框和模糊截图，替换为完整页面截图，保证菜单、标题、按钮和结果区域可见。"],
        ],
        widths=[Cm(4), Cm(12)],
    )

    doc.add_heading("二、现有文档问题诊断", level=1)
    add_body_paragraph(doc, "根据当前操作手册内容，建议重点修正以下问题：")
    add_table(
        doc,
        ["问题类型", "表现", "修改方式"],
        [
            ["截图缺失", "原第7页“图1 观测记录界面”“图2 多源数据便捷上传”位置仅剩图框和文字。", "重新插入完整清晰截图，并检查图片是否嵌入文档而非外链。"],
            ["截图链条不完整", "部分功能只展示结果页面，没有展示入口、填写、提交等中间步骤。", "每个核心功能补齐“入口图、操作图、结果图”。"],
            ["文字偏概述", "多处写成系统介绍，如“系统支持多源数据采集”。", "改成操作说明，如“用户点击新增按钮，填写字段，保存后返回列表”。"],
            ["管理员功能展开不足", "数据审核、SHP上传、样线绘制、权限和积分商城说明较少。", "将管理员端单独成章，每个功能配步骤和截图。"],
            ["图号与正文风险", "补图后容易出现图号跳号、正文引用不一致。", "最终统一更新图号、目录和正文“如图X”引用。"],
        ],
        widths=[Cm(3.2), Cm(6.2), Cm(6.6)],
    )

    doc.add_heading("三、建议重构后的文档目录", level=1)
    add_body_paragraph(doc, "建议将原手册重构为以下章节结构。章节不必追求过多，但必须覆盖完整使用流程和全部主要功能。")
    add_table(
        doc,
        ["章节", "建议标题", "修改目的"],
        [
            ["封面", "黄河生态方舟系统 V1.0 操作手册", "统一软件名称、版本号、开发单位等基础信息。"],
            ["第一章", "系统概述", "简要说明软件用途、角色、功能范围，避免大段背景占用篇幅。"],
            ["第二章", "系统运行环境", "说明服务端、客户端、浏览器和网络环境。"],
            ["第三章", "系统登录与首页", "补齐完整使用流程的起点。"],
            ["第四章", "用户端功能操作", "说明巡护人员或普通用户如何采集、上传、查询和查看地图数据。"],
            ["第五章", "管理员端功能操作", "说明管理员如何审核、维护、导入空间数据、管理用户和积分商城。"],
            ["第六章", "常见操作说明", "说明保存成功、查询无结果、图片上传失败、权限不足等常见场景。"],
            ["第七章", "附录：主要功能清单", "用表格明确列出功能覆盖情况，回应“全部功能”要求。"],
        ],
        widths=[Cm(2.5), Cm(5.5), Cm(8)],
    )

    doc.add_heading("四、逐章修改策略", level=1)

    doc.add_heading("4.1 封面", level=2)
    add_body_paragraph(doc, "封面应保持正式、简洁，并与软著申请表中的软件名称和版本号一致。建议删除多余表述，统一为以下信息：")
    add_bullets(
        doc,
        [
            "软件全称：黄河生态方舟系统",
            "软件简称：黄河生态方舟",
            "版本号：V1.0",
            "文档名称：操作手册",
            "开发单位：郑州大学",
            "编写人：周轩",
            "编写日期：2026年1月28日",
        ],
    )
    add_body_paragraph(doc, "注意事项：全文软件名称、简称、版本号必须保持一致，不要交替使用“生态方舟平台”“黄河生态平台”等其他名称。")

    doc.add_heading("4.2 目录", level=2)
    add_body_paragraph(doc, "目录应根据最终章节自动更新。补正完成后，应右键更新目录，选择“更新整个目录”，确保标题和页码与正文一致。")
    add_bullets(
        doc,
        [
            "新增“系统登录与首页”章节，体现软件使用起点。",
            "用户端和管理员端分章，避免两类角色操作混杂。",
            "目录中的标题应与正文标题完全一致。",
            "最终提交前检查页码是否正确，避免旧目录残留。",
        ],
    )

    doc.add_heading("4.3 第一章 系统概述", level=2)
    add_body_paragraph(doc, "第一章不宜写得过长。原有背景、简介、目标用户、主要功能可保留但应压缩，重点改为角色和功能范围说明。")
    add_table(
        doc,
        ["小节", "应如何修改", "建议内容"],
        [
            ["1.1 软件基本信息", "增加基础信息表。", "软件名称、简称、版本号、软件类型、使用对象、主要功能。"],
            ["1.2 系统功能概述", "由长段落改为分条说明。", "用户端用于采集、上传、查询、地图查看；管理员端用于审核、维护、统计、权限和积分管理。"],
            ["1.3 用户角色说明", "增加角色权限表。", "普通用户、巡护人员、管理员、系统管理员分别列权限。"],
            ["1.4 使用流程总览", "新增流程说明。", "登录系统 → 进入首页 → 选择模块 → 操作数据 → 保存/查询 → 查看结果。"],
        ],
        widths=[Cm(4), Cm(5.2), Cm(6.8)],
    )
    add_caption(doc, "表1 第一章修改策略")

    doc.add_heading("4.4 第二章 系统运行环境", level=2)
    add_body_paragraph(doc, "第二章可保留现有技术环境，但建议删减过细或可能敏感的信息，例如服务器实例号、账号、内网地址、密钥等。")
    add_bullets(
        doc,
        [
            "服务端环境：说明操作系统、Django、PostgreSQL/PostGIS、Nginx、Gunicorn等。",
            "客户端环境：说明Chrome、Edge、Safari等浏览器及移动端访问方式。",
            "网络环境：说明系统通过网络访问，地图功能需在线地图服务支持。",
            "安全处理：截图中如出现真实IP、域名、账号、手机号、密钥，应打码。",
        ],
    )

    doc.add_heading("4.5 第三章 系统登录与首页", level=2)
    add_body_paragraph(doc, "第三章建议新增，是完整使用流程的开头。若没有登录与首页说明，审查人员会认为文档不是完整操作手册。")
    add_table(
        doc,
        ["小节", "文字应写什么", "必须配套截图"],
        [
            ["3.1 系统访问", "用户在浏览器输入系统访问地址，进入登录页面。", "系统登录页面。"],
            ["3.2 用户登录", "输入用户名、密码，点击登录，验证通过后进入首页。", "输入账号密码、登录成功首页。"],
            ["3.3 首页功能区说明", "说明顶部导航、左侧菜单、地图区域、统计区域、数据列表等。", "首页完整截图，能看到菜单和主要区域。"],
            ["3.4 退出登录", "点击用户头像或退出按钮，返回登录页面。", "退出入口或退出后页面。"],
        ],
        widths=[Cm(3.5), Cm(7), Cm(5.5)],
    )
    add_body_paragraph(doc, "截图要求：登录页截图应完整显示系统名称、账号框、密码框、登录按钮；首页截图应完整显示左侧菜单和主要功能区域。")

    doc.add_heading("4.6 第四章 用户端功能操作", level=2)
    add_body_paragraph(doc, "第四章是补正重点之一。建议将原第三章“系统主要功能（用户）”整体改写为连续操作流程，而不是功能概述。")
    add_table(
        doc,
        ["功能", "应补写的操作步骤", "建议截图"],
        [
            ["观测记录列表", "登录后点击“观测记录”，查看物种名称、时间、地点、状态等列表信息。", "观测记录列表完整页面。"],
            ["新增观测记录", "点击新增，填写物种、时间、地点、数量、环境描述，保存。", "新增入口、表单填写、保存成功列表。"],
            ["上传图片/视频", "在新增或编辑页面点击上传，选择文件，上传完成后保存。", "选择文件、上传完成、图片关联结果。"],
            ["数据查询", "选择时间、物种、点位等查询条件，点击查询，查看结果。", "设置条件、查询结果列表。"],
            ["地图点位详情", "进入地图，点击点位，查看弹窗详情。", "地图点位分布、点位详情弹窗。"],
            ["图层切换", "点击图层控件，切换电子地图、卫星影像、监测点、样线和观测记录。", "图层面板、各图层展示效果。"],
            ["监测点位查看", "勾选监测点位图层，点击点位查看名称、编号、位置。", "监测点位展示、点位详情。"],
            ["监测样线查看", "勾选监测样线图层，通过缩放和拖动查看线路范围。", "样线地图展示。"],
        ],
        widths=[Cm(3.2), Cm(7.3), Cm(5.5)],
    )
    add_body_paragraph(doc, "特别注意：原第7页缺失的“观测记录界面”和“多源数据便捷上传”截图必须重新插入，且应分别对应“观测记录列表”和“新增/上传观测记录”两个真实操作步骤。")

    doc.add_heading("4.7 第五章 管理员端功能操作", level=2)
    add_body_paragraph(doc, "第五章应单独展开管理员功能，不能只用一两段概述。管理员端是证明软件具备完整后台管理能力的重要部分。")
    add_table(
        doc,
        ["功能", "应补写的操作步骤", "建议截图"],
        [
            ["管理员后台首页", "管理员账号登录后进入后台，左侧显示数据管理、空间数据、用户管理、积分商城等菜单。", "管理员后台首页。"],
            ["观测数据审核", "进入观测记录管理，查看记录详情，核对数据和图片，选择通过或退回。", "记录列表、详情页、审核操作、状态更新。"],
            ["数据修改", "点击编辑，修改错误字段，保存后返回列表。", "编辑入口、编辑表单、保存结果。"],
            ["数据删除", "点击删除，确认后删除无效数据。", "删除确认、删除后列表。"],
            ["SHP空间数据上传", "进入load_shp或空间数据管理模块，选择文件，填写图层信息，提交上传。", "模块入口、文件选择、提交页面、上传成功、地图展示。"],
            ["手绘监测样线", "进入样线管理，点击绘制，在地图上绘制线路，填写属性并保存。", "绘制入口、地图绘制、属性填写、保存结果。"],
            ["驾驶舱统计", "进入驾驶舱查看总量、趋势、分类统计和空间分布。", "驾驶舱总览、统计图表。"],
            ["用户管理", "查看用户列表，新增、编辑、禁用用户或重置密码。", "用户列表、新增用户、编辑用户。"],
            ["权限管理", "为不同角色配置菜单和数据权限。", "角色管理、权限配置。"],
            ["积分商城管理", "维护商品名称、积分、库存、图片和兑换状态，查看兑换记录。", "商品列表、新增商品、编辑商品、兑换记录。"],
        ],
        widths=[Cm(3.2), Cm(7.3), Cm(5.5)],
    )
    add_body_paragraph(doc, "如实际系统没有独立权限管理页面，不要虚构功能，可改写为“用户角色设置”或“用户权限字段设置”，以真实界面为准。")

    doc.add_heading("4.8 第六章 常见操作说明", level=2)
    add_body_paragraph(doc, "第六章用于增强文档完整性，说明用户在实际使用中可能遇到的常见结果或提示。")
    add_bullets(
        doc,
        [
            "数据保存成功：说明保存后记录写入数据库，并可在列表或地图中查看。",
            "查询无结果：说明条件过窄或暂无相关数据，可调整时间、物种、点位等条件。",
            "图片上传失败：说明文件格式、大小、网络状态可能影响上传。",
            "地图加载异常：说明需检查网络连接或在线地图服务访问状态。",
            "权限不足：说明不同角色可访问功能不同，需联系管理员调整权限。",
        ],
    )

    doc.add_heading("4.9 第七章 附录：主要功能清单", level=2)
    add_body_paragraph(doc, "附录建议放一张功能覆盖表，直接回应补正意见中“全部功能”的要求。")
    add_table(
        doc,
        ["序号", "模块", "功能名称", "是否应有截图"],
        [
            ["1", "登录模块", "用户登录", "是"],
            ["2", "首页模块", "首页/驾驶舱入口", "是"],
            ["3", "用户端", "新增观测记录", "是"],
            ["4", "用户端", "上传图片/视频", "是"],
            ["5", "用户端", "数据查询", "是"],
            ["6", "用户端", "地图点位查看", "是"],
            ["7", "用户端", "图层切换", "是"],
            ["8", "管理员端", "观测数据审核", "是"],
            ["9", "管理员端", "数据修改/删除", "是"],
            ["10", "管理员端", "SHP上传", "是"],
            ["11", "管理员端", "样线绘制", "是"],
            ["12", "管理员端", "用户和权限管理", "是"],
            ["13", "管理员端", "积分商城管理", "是"],
        ],
        widths=[Cm(2), Cm(4), Cm(6), Cm(4)],
    )

    doc.add_heading("五、统一写法模板", level=1)
    add_body_paragraph(doc, "为保证全文连贯，建议每个功能均采用以下统一结构：")
    add_table(
        doc,
        ["组成部分", "写法要求", "示例"],
        [
            ["功能说明", "说明该功能用于解决什么问题。", "新增观测记录用于巡护人员录入现场生态监测数据。"],
            ["操作入口", "说明从哪个菜单或按钮进入。", "用户登录后，在左侧菜单点击“观测记录”。"],
            ["操作步骤", "按第一步、第二步、第三步描述。", "点击“新增”按钮，填写物种名称、观测时间、地点等字段。"],
            ["操作结果", "说明完成后页面变化。", "保存成功后，记录显示在列表中，并可在地图中查看点位。"],
            ["配套截图", "列出该功能需要的截图。", "观测记录列表、新增表单、上传图片、保存结果。"],
        ],
        widths=[Cm(3), Cm(5.5), Cm(7.5)],
    )

    doc.add_heading("六、重点段落改写示例", level=1)
    add_body_paragraph(doc, "原说明常见问题是偏“介绍”，修改后应变成“操作”。可参考以下示例。")
    add_table(
        doc,
        ["原写法", "建议改写"],
        [
            [
                "系统支持多源数据便捷采集与上传。",
                "用户进入“观测记录”页面后，点击“新增”按钮，填写物种名称、观测时间、观测地点、数量、环境描述等信息，并上传现场图片。点击“保存”后，系统将数据写入观测记录列表，并在地图页面展示对应点位。",
            ],
            [
                "系统提供按时间条件查询，方便用户筛选数据记录。",
                "用户在观测记录列表上方选择开始时间和结束时间，点击“查询”按钮，系统刷新列表并显示该时间范围内的观测记录。若需要恢复全部数据，可清空时间条件后重新查询。",
            ],
            [
                "管理员可对所有用户上传的观测数据进行审核和修正。",
                "管理员进入“观测记录管理”页面，点击待审核记录后的“查看”按钮，核对物种、时间、地点和现场图片。确认无误后点击“审核通过”；如发现错误，可点击“编辑”修改字段或退回记录。",
            ],
        ],
        widths=[Cm(6), Cm(10)],
    )

    doc.add_heading("七、截图补充清单", level=1)
    add_body_paragraph(doc, "建议按下表逐项补齐截图。截图不需要过度美化，但必须清晰、完整、能对应正文操作。")
    screenshot_rows = [
        ["1", "登录页面", "完整显示系统名称、账号框、密码框、登录按钮。"],
        ["2", "登录成功首页", "显示顶部导航、左侧菜单、主要功能区。"],
        ["3", "观测记录列表", "替换原第7页缺失截图。"],
        ["4", "点击新增观测记录", "显示新增按钮所在位置。"],
        ["5", "新增观测记录表单", "显示主要字段和保存按钮。"],
        ["6", "上传图片/视频", "显示文件选择或上传完成状态。"],
        ["7", "保存成功后的记录列表", "显示新增记录出现在列表中。"],
        ["8", "时间查询条件", "显示开始时间、结束时间和查询按钮。"],
        ["9", "查询结果", "显示筛选后的列表。"],
        ["10", "地图点位详情", "显示地图点位和详情弹窗。"],
        ["11", "图层切换面板", "显示电子地图、卫星影像、监测点、样线、观测记录等选项。"],
        ["12", "监测点位图层", "显示点位在地图中的空间分布。"],
        ["13", "监测样线图层", "显示样线在地图中的空间走向。"],
        ["14", "管理员后台首页", "显示后台菜单和主要区域。"],
        ["15", "观测数据审核", "显示审核入口、详情页和审核结果。"],
        ["16", "数据修改/删除", "显示编辑表单和删除确认。"],
        ["17", "SHP上传", "显示文件选择、提交和上传成功结果。"],
        ["18", "手绘样线", "显示地图绘制、属性填写和保存结果。"],
        ["19", "驾驶舱统计", "显示数据概览和统计图表。"],
        ["20", "用户管理", "显示用户列表、新增或编辑用户页面。"],
        ["21", "权限管理", "显示角色或权限配置页面。"],
        ["22", "积分商城管理", "显示商品列表、编辑页面和兑换记录。"],
    ]
    add_table(doc, ["序号", "截图名称", "截图要求"], screenshot_rows, widths=[Cm(1.7), Cm(4.5), Cm(9.8)])

    doc.add_heading("八、截图规范", level=1)
    add_bullets(
        doc,
        [
            "浏览器缩放建议设置为100%，截图分辨率尽量使用1920×1080或更高。",
            "每张截图应显示完整系统页面，至少包含页面标题、左侧菜单、主要按钮和内容区域。",
            "一页建议放1张大图，最多放2张，避免图片缩得过小导致文字不可辨认。",
            "所有空白图框、损坏图片、严重裁切图片、模糊图片必须删除或替换。",
            "图片下方统一使用“图4-3 新增观测记录表单”这类图题。",
            "正文中的“如图X”必须与实际图题一致。",
            "截图中涉及真实账号、密码、手机号、密钥、服务器地址等信息时，应进行遮挡或打码。",
        ]
    )

    doc.add_heading("九、补正提交说明", level=1)
    add_body_paragraph(doc, "完成修改后，建议将文件另存为“黄河生态方舟系统V1.0操作手册_补正版.docx”。如补正系统允许上传PDF，建议同时导出PDF，以降低版式错乱风险。")
    add_body_paragraph(doc, "补正说明可使用以下文字：")
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.right_indent = Cm(0.8)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(
        "已根据补正通知要求，对软件操作手册进行修改完善：补充了系统登录、用户端数据采集、数据查询、地图展示、管理员数据审核、SHP上传、样线绘制、用户权限管理、积分商城管理等主要功能的详细操作流程；重新补充并替换了不完整截图，确保截图清晰、完整、与正文说明对应。"
    )
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(10.5)
    r.italic = True

    doc.add_heading("十、最终自查清单", level=1)
    checklist = [
        "第7页原缺失截图已替换。",
        "所有图片均能正常显示，没有空白图框。",
        "用户端主要功能均有操作步骤和截图。",
        "管理员端主要功能均有操作步骤和截图。",
        "每个核心功能至少包含入口、操作、结果三个环节。",
        "图号连续，正文引用图号正确。",
        "目录已更新，页码正确。",
        "软件名称、简称、版本号与软著申请信息一致。",
        "截图中无明文密码、密钥、敏感账号或服务器敏感信息。",
        "文件可正常打开，另存为补正版后再上传。",
    ]
    add_table(doc, ["序号", "检查事项", "完成情况"], [[i + 1, item, "□"] for i, item in enumerate(checklist)], widths=[Cm(2), Cm(11), Cm(3)])

    doc.add_heading("十一、建议执行顺序", level=1)
    add_numbered(
        doc,
        [
            "先按新目录重排文档结构，新增“系统登录与首页”和“常见操作说明”。",
            "重新截取登录、首页、用户端、管理员端所有关键页面。",
            "优先替换原第7页缺失截图，并删除所有空白图框。",
            "逐个功能按“功能说明、操作入口、操作步骤、操作结果、配套截图”改写。",
            "补齐管理员端审核、SHP上传、样线绘制、用户权限、积分商城等功能。",
            "统一图题格式和正文引用。",
            "更新目录，检查页码、图号、错别字和敏感信息。",
            "另存补正版，并根据平台要求上传DOCX或PDF。",
        ],
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
