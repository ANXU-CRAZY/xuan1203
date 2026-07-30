#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""生成计算机设计大赛省赛设计与开发文档（Word格式）"""
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ========== 全局样式设置 ==========
style = doc.styles['Normal']
style.font.name = '宋体'
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
style.paragraph_format.line_spacing = 1.5

# 标题样式
for i in range(1, 4):
    h = doc.styles[f'Heading {i}']
    h.font.name = '黑体'
    h.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    h.font.color.rgb = RGBColor(0, 0, 0)

doc.styles['Heading 1'].font.size = Pt(22)
doc.styles['Heading 2'].font.size = Pt(16)
doc.styles['Heading 3'].font.size = Pt(14)


def add_cover():
    """封面页"""
    for _ in range(4):
        doc.add_paragraph()
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('中国大学生计算机设计大赛')
    run.font.size = Pt(26)
    run.font.name = '黑体'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    run.bold = True

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('设计与开发文档')
    run.font.size = Pt(22)
    run.font.name = '黑体'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

    for _ in range(3):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('作品名称：黄河生态方舟')
    run.font.size = Pt(18)
    run.font.name = '宋体'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('——面向黄河湿地鸟类保护的智能监测与公众协同平台')
    run.font.size = Pt(14)
    run.font.name = '宋体'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    for _ in range(4):
        doc.add_paragraph()

    info_lines = [
        '参赛类别：软件应用与开发 — Web 应用与开发',
        '参赛学校：____________________',
        '团队成员：____________________',
        '指导教师：____________________',
        '完成日期：2026 年 5 月',
    ]
    for line in info_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        run.font.size = Pt(14)

    doc.add_page_break()


def add_toc():
    """目录页"""
    doc.add_heading('目  录', level=1)
    toc_items = [
        ('一、项目概述', 1), ('二、需求分析', 3), ('三、系统总体设计', 5),
        ('四、详细设计与实现', 8), ('五、数据库设计', 14), ('六、系统测试', 16),
        ('七、部署与运维', 18), ('八、创新点与应用价值', 19), ('九、总结与展望', 21),
    ]
    for title, page in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(15))
        run = p.add_run(f'{title}')
        run.font.size = Pt(14)
    doc.add_page_break()


def add_chapter1():
    """第一章 项目概述"""
    doc.add_heading('一、项目概述', level=1)

    doc.add_heading('1.1 项目背景', level=2)
    doc.add_paragraph(
        '黄河流域生态保护和高质量发展是重大国家战略。黄河中下游湿地是东亚—澳大利西亚候鸟迁飞路线的重要节点，'
        '每年有数百种、数十万只候鸟在此停歇、越冬和繁殖。然而，传统湿地鸟类监测面临数据分散、公众参与门槛高、'
        '物种知识传播效率低、管理端审核与展示链路割裂等问题。'
    )
    doc.add_paragraph(
        '在此背景下，本项目设计并实现了"黄河生态方舟"——一套面向黄河湿地鸟类保护与公众协同巡护的 Web 综合平台，'
        '构建"观测采集—数据治理—智能识别—地图分析—科普传播—互动参与"的一体化应用系统，'
        '为黄河湿地鸟类保护提供数字化、可视化、可扩展的解决方案。'
    )

    doc.add_heading('1.2 项目目标', level=2)
    goals = [
        '建立统一的黄河湿地鸟类观测数据管理平台，实现数据从采集到应用的完整闭环。',
        '通过地图可视化与热力图分析，直观呈现鸟类活动的时空分布规律。',
        '集成水鸟图像识别模型，降低公众参与物种辨识的门槛。',
        '构建物种百科、图库与科普文章体系，提升公众生态保护意识。',
        '提供游戏化互动模块，让生态教育从静态文字转化为可体验的交互过程。',
    ]
    for g in goals:
        doc.add_paragraph(g, style='List Bullet')

    doc.add_heading('1.3 参赛定位', level=2)
    doc.add_paragraph(
        '本项目主报"软件应用与开发—Web 应用与开发"赛道。项目主体为 B/S 架构的生态监测与公众参与平台，'
        '包含地图可视化、数据管理、物种百科、观鸟记录、图像识别、科普图库、游戏化互动、后台审核与部署上线能力，'
        '符合该赛道对"运行在网络、数据库系统之上的软件，提供信息管理、信息服务、算法应用等功能"的要求。'
    )
    doc.add_page_break()


def add_chapter2():
    """第二章 需求分析"""
    doc.add_heading('二、需求分析', level=1)

    doc.add_heading('2.1 用户角色分析', level=2)
    roles = [
        ('游客/公众用户', '浏览生态大屏、物种百科、图库、科普文章和互动游戏，了解黄河湿地鸟类资源。'),
        ('观鸟爱好者/巡护人员', '通过上报入口提交观鸟记录（物种、点位、时间、数量、图片），参与数据采集。'),
        ('管理员', '通过后台管理物种信息、区域、样线、观测记录、图库和识别历史，审核公众上报数据。'),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    hdr[0].text = '用户角色'
    hdr[1].text = '核心需求'
    for role, desc in roles:
        row = table.add_row().cells
        row[0].text = role
        row[1].text = desc

    doc.add_heading('2.2 功能需求', level=2)
    doc.add_heading('2.2.1 数据采集与管理', level=3)
    items = [
        '支持观鸟记录的在线上报，包含物种、点位、时间、数量、图片和描述。',
        '管理员可对上报数据进行审核（通过/驳回），保证数据质量。',
        '支持 CSV 批量导入历史观测数据。',
    ]
    for item in items:
        doc.add_paragraph(item, style='List Bullet')

    doc.add_heading('2.2.2 地图可视化与空间分析', level=3)
    items = [
        '基于天地图底图展示湿地区域、监测点位、监测样线和观鸟记录。',
        '支持图层切换、时间范围筛选和保护等级筛选。',
        '支持区域圈选生成鸟类数量热力图。',
    ]
    for item in items:
        doc.add_paragraph(item, style='List Bullet')

    doc.add_heading('2.2.3 水鸟图像识别', level=3)
    items = [
        '用户上传鸟类图片，系统自动检测鸟体区域并进行物种分类。',
        '支持检测阈值和分类阈值设置，适应不同图片质量。',
    ]
    for item in items:
        doc.add_paragraph(item, style='List Bullet')

    doc.add_heading('2.2.4 物种百科与科普传播', level=3)
    items = [
        '展示物种中文名、拉丁名、目科、保护等级、习性和图库。',
        '自动生成物种科普文章，支持搜索和分类浏览。',
        '提供多个湿地主题互动小游戏，提升公众参与度。',
    ]
    for item in items:
        doc.add_paragraph(item, style='List Bullet')

    doc.add_heading('2.3 非功能需求', level=2)
    items = [
        '性能：首页地图加载时间不超过 5 秒，API 响应时间不超过 3 秒。',
        '安全：用户认证基于 Token，敏感操作需登录，API Key 不暴露于前端。',
        '可用性：系统部署于云服务器，支持 7×24 小时访问。',
        '兼容性：支持主流浏览器（Chrome、Edge、Firefox）和移动端自适应。',
    ]
    for item in items:
        doc.add_paragraph(item, style='List Bullet')
    doc.add_page_break()


add_cover()
add_toc()
add_chapter1()
add_chapter2()

# 保存临时文件，后续追加内容
doc.save('design_doc_temp.docx')
print('第1-2章已生成')
