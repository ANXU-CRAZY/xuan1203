#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""生成设计文档第3-9章"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document('design_doc_temp.docx')


def add_chapter3():
    """第三章 系统总体设计"""
    doc.add_heading('三、系统总体设计', level=1)

    doc.add_heading('3.1 系统架构设计', level=2)
    doc.add_paragraph(
        '本系统采用经典的 B/S（Browser/Server）架构，分为前端展示层、后端业务层和数据持久层三个层次。'
        '前端负责用户交互与数据可视化，后端提供 RESTful API 接口和业务逻辑处理，数据层使用 PostgreSQL 结合 PostGIS 扩展存储空间数据。'
    )
    doc.add_paragraph('系统架构分层如下：')
    layers = [
        '表现层：HTML5 + CSS3 + JavaScript + Leaflet 地图 + Chart.js 图表',
        '接口层：Django REST Framework 提供 RESTful API',
        '业务层：Django 应用（app_monitor、bird_recognition）',
        '数据层：PostgreSQL + PostGIS 空间数据库',
        'AI 层：YOLOv8 检测模型 + 分类模型',
        '部署层：Nginx + Gunicorn + 阿里云 ECS',
    ]
    for layer in layers:
        doc.add_paragraph(layer, style='List Bullet')

    doc.add_heading('3.2 技术选型', level=2)
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    hdr[0].text = '层次'
    hdr[1].text = '技术/工具'
    hdr[2].text = '选型理由'
    techs = [
        ('前端', 'HTML5/CSS3/JavaScript', '无需编译，快速迭代，兼容性好'),
        ('地图引擎', 'Leaflet + 天地图', '轻量开源，国内底图合规'),
        ('后端框架', 'Django 5.2 + DRF', '成熟稳定，ORM 强大，生态丰富'),
        ('数据库', 'PostgreSQL 16 + PostGIS', '支持空间查询，性能优秀'),
        ('AI 模型', 'YOLOv8 (Ultralytics)', '检测精度高，推理速度快'),
        ('部署', 'Nginx + Gunicorn', '高并发，稳定可靠'),
        ('版本管理', 'Git + GitHub', '团队协作，代码追溯'),
    ]
    for layer, tech, reason in techs:
        row = table.add_row().cells
        row[0].text = layer
        row[1].text = tech
        row[2].text = reason

    doc.add_heading('3.3 功能模块划分', level=2)
    doc.add_paragraph('系统包含以下核心功能模块：')
    modules = [
        '生态大屏模块：地图展示、图层控制、时间筛选、统计面板、热力图圈选。',
        '观测记录模块：数据上报、审核管理、CSV 导入、地图标注。',
        '水鸟识别模块：图片上传、鸟体检测、物种分类、阈值设置。',
        '物种百科模块：物种信息展示、图库管理、科普文章生成。',
        '用户系统模块：注册登录、积分管理、个人中心、头像上传。',
        '互动游戏模块：鸟类猜图、迁徙模拟、湿地修复、湿地侦探、浮岛生态。',
        '智能助手模块：DeepSeek 大模型驱动的生态问答助手。',
    ]
    for m in modules:
        doc.add_paragraph(m, style='List Bullet')
    doc.add_page_break()


def add_chapter4():
    """第四章 详细设计与实现"""
    doc.add_heading('四、详细设计与实现', level=1)

    doc.add_heading('4.1 生态大屏模块', level=2)
    doc.add_paragraph(
        '生态大屏是系统的核心展示页面，基于 Leaflet 地图引擎和天地图底图构建。'
        '页面集成了监测点位、监测样线、观鸟记录三类空间数据图层，支持图层独立开关、时间范围筛选和保护等级筛选。'
    )
    doc.add_heading('4.1.1 地图图层管理', level=3)
    doc.add_paragraph(
        '系统提供电子地图、卫星影像两种底图，以及行政区界、监测点位、监测样线、观鸟记录四个叠加图层。'
        '用户可通过左侧图层控制面板自由切换，各图层数据通过 AJAX 异步加载，不阻塞页面渲染。'
    )
    doc.add_heading('4.1.2 热力图圈选分析', level=3)
    doc.add_paragraph(
        '用户在地图上拖拽绘制矩形区域后，系统自动统计该范围内的观鸟记录数量，'
        '按网格划分生成鸟类数量热力图。热力图使用 Canvas 渲染，支持颜色梯度映射，'
        '直观呈现鸟类活动热点区域，辅助巡护决策。'
    )
    doc.add_heading('4.1.3 观测记录缓存优化', level=3)
    doc.add_paragraph(
        '系统建立了 MapObservationCache 缓存表，将 93000+ 条审核通过的观测记录预计算并缓存，'
        '避免每次请求都进行多表关联查询。API 支持视口范围过滤（bbox 参数），'
        '仅返回当前地图可视区域内的记录，配合 5 分钟服务端缓存，显著提升加载速度。'
    )

    doc.add_heading('4.2 水鸟图像识别模块', level=2)
    doc.add_paragraph(
        '水鸟识别模块采用两阶段流水线架构：第一阶段使用 YOLOv8 检测模型定位图片中的鸟体区域，'
        '第二阶段使用分类模型对裁剪后的鸟体图像进行物种判断。'
    )
    doc.add_heading('4.2.1 检测阶段', level=3)
    doc.add_paragraph(
        '检测模型（detector.pt）基于 YOLOv8 架构训练，输入为用户上传的原始图片，'
        '输出为图片中所有鸟体的边界框坐标和置信度。用户可设置检测阈值（默认 0.5），'
        '低于阈值的检测结果将被过滤。'
    )
    doc.add_heading('4.2.2 分类阶段', level=3)
    doc.add_paragraph(
        '分类模型（classifier.pt）对检测到的每个鸟体区域进行物种分类，'
        '输出 Top-5 候选物种及其置信度。用户可设置分类阈值（默认 0.3），'
        '仅展示高于阈值的分类结果。识别结果同时记录到 AIDetectionResult 表中，便于后续分析。'
    )

    doc.add_heading('4.3 物种百科与图库模块', level=2)
    doc.add_paragraph(
        '物种百科模块管理 431 个鸟类物种的结构化信息，包括中文名、拉丁名、目、科、保护等级和分布习性。'
        '每个物种关联图库图片，系统通过多级图片兜底策略确保展示效果：'
    )
    items = [
        '优先级 1：数据库 SpeciesInfo.cover_image 字段（管理员上传）。',
        '优先级 2：SpeciesImage 表中的精选图片（is_featured=True）。',
        '优先级 3：SpeciesImage 表中的第一张图片。',
        '优先级 4：前端模板 SPECIES_IMG 映射（135 种 Wikimedia Commons 直链）。',
        '优先级 5：基于拉丁名自动生成的 Wikimedia Commons 查询 URL。',
    ]
    for item in items:
        doc.add_paragraph(item, style='List Bullet')

    doc.add_heading('4.4 用户系统与积分模块', level=2)
    doc.add_paragraph(
        '系统基于 Django 内置认证框架和 Token 认证实现用户管理。'
        '用户注册后自动创建积分档案（UserProfile），每次成功上报观鸟记录奖励 10 积分。'
        '积分可在积分商城兑换虚拟商品，形成"参与—贡献—激励"的正向循环。'
    )

    doc.add_heading('4.5 智能助手模块', level=2)
    doc.add_paragraph(
        '系统集成 DeepSeek 大语言模型作为生态问答助手。助手以当前地图筛选条件、'
        '统计数据、热门物种和热门区域作为上下文，生成自然语言的监测摘要和巡护建议。'
        '当 API 不可用时，系统自动降级为本地数据摘要模式，确保功能可用性。'
    )

    doc.add_heading('4.6 互动游戏模块', level=2)
    doc.add_paragraph(
        '为提升公众参与度和科普传播效果，系统提供 6 个湿地主题互动小游戏：'
    )
    games = [
        '鸟类猜图：展示鸟类图片，用户猜测物种名称，学习鸟类外形特征。',
        '鸟类跑酷：横版跑酷游戏，躲避障碍收集鸟类知识卡片。',
        '迁徙模拟：模拟候鸟迁徙路线，了解迁飞通道和停歇地。',
        '湿地修复：策略类游戏，通过种植、清理等操作恢复湿地生态。',
        '湿地侦探：寻找隐藏在湿地场景中的鸟类，锻炼观察力。',
        '浮岛生态：经营浮岛生态系统，维持生物多样性平衡。',
    ]
    for g in games:
        doc.add_paragraph(g, style='List Bullet')
    doc.add_page_break()


def add_chapter5():
    """第五章 数据库设计"""
    doc.add_heading('五、数据库设计', level=1)

    doc.add_heading('5.1 数据库选型', level=2)
    doc.add_paragraph(
        '系统使用 PostgreSQL 16 作为主数据库，配合 PostGIS 扩展支持空间数据存储与查询。'
        'PostGIS 提供了 Point、MultiLineString 等空间数据类型和 ST_DWithin、ST_Intersects 等空间函数，'
        '满足地图展示和空间分析的需求。'
    )

    doc.add_heading('5.2 核心数据表设计', level=2)
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    hdr[0].text = '表名'
    hdr[1].text = '说明'
    hdr[2].text = '核心字段'
    hdr[3].text = '记录数'
    tables_data = [
        ('SpeciesInfo', '物种信息', 'name_cn, name_latin, order, family, protection_level', '431'),
        ('WetlandZone', '监测点位', 'name, latitude, longitude, location(Point)', '69'),
        ('MonitoringRoute', '监测样线', 'name, path_geom(MultiLineString)', '12'),
        ('ObservationRecord', '观测记录', 'species, zone, count, observation_time, status', '93000+'),
        ('MapObservationCache', '地图缓存', '预计算的观测记录快照，含坐标和物种信息', '93000+'),
        ('SpeciesImage', '物种图库', 'species, image, image_url, source, is_featured', '400+'),
        ('AIDetectionResult', 'AI识别记录', 'image, species_name, confidence, model_version', '动态'),
        ('UserProfile', '用户档案', 'user, score, avatar', '动态'),
        ('Product', '积分商品', 'name, price, stock, description', '若干'),
    ]
    for name, desc, fields, count in tables_data:
        row = table.add_row().cells
        row[0].text = name
        row[1].text = desc
        row[2].text = fields
        row[3].text = count

    doc.add_heading('5.3 空间数据设计', level=2)
    doc.add_paragraph(
        '观测记录通过 PostGIS 的 Point 类型存储精确坐标（SRID=4326，WGS84 坐标系）。'
        '监测样线使用 MultiLineString 类型存储路径几何。系统支持基于空间范围的查询过滤（bbox），'
        '以及基于距离的附近预警查询（ST_DWithin）。'
    )

    doc.add_heading('5.4 缓存策略', level=2)
    doc.add_paragraph(
        '为优化首页地图加载性能，系统设计了 MapObservationCache 缓存表，'
        '将多表关联查询的结果预计算存储。缓存表包含观测记录的坐标、物种名称、保护等级等展示所需字段，'
        '避免每次请求都进行复杂的 JOIN 操作。同时配合 Django 缓存框架实现 5 分钟的 API 响应缓存。'
    )
    doc.add_page_break()


def add_chapter6():
    """第六章 系统测试"""
    doc.add_heading('六、系统测试', level=1)

    doc.add_heading('6.1 测试环境', level=2)
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    hdr[0].text = '项目'
    hdr[1].text = '配置'
    envs = [
        ('开发环境', 'Windows 11, Python 3.11, Django 5.2.9, PostgreSQL 16'),
        ('服务器', '阿里云 ECS 2核4G, Ubuntu 24.04, Nginx 1.24, Gunicorn'),
        ('数据库', 'PostgreSQL 16 + PostGIS 3.4'),
        ('浏览器', 'Chrome 148, Edge 148, Firefox 最新版'),
    ]
    for item, config in envs:
        row = table.add_row().cells
        row[0].text = item
        row[1].text = config

    doc.add_heading('6.2 功能测试', level=2)
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    hdr[0].text = '测试项'
    hdr[1].text = '测试内容'
    hdr[2].text = '预期结果'
    hdr[3].text = '实际结果'
    tests = [
        ('地图加载', '首页打开后地图正常渲染', '天地图底图和图层正常显示', '通过'),
        ('观测记录', '筛选时间范围后查看记录', '仅显示范围内的记录点', '通过'),
        ('热力图', '圈选区域生成热力图', '正确统计并渲染热力色块', '通过'),
        ('水鸟识别', '上传鸟类图片进行识别', '返回检测框和分类结果', '通过'),
        ('物种百科', '搜索物种查看详情', '展示物种信息和图片', '通过'),
        ('用户注册', '填写信息注册新用户', '注册成功并返回 Token', '通过'),
        ('数据上报', '登录后提交观鸟记录', '记录保存并积分增加', '通过'),
        ('后台审核', '管理员审核上报记录', '状态变更为已通过', '通过'),
    ]
    for t1, t2, t3, t4 in tests:
        row = table.add_row().cells
        row[0].text = t1
        row[1].text = t2
        row[2].text = t3
        row[3].text = t4

    doc.add_heading('6.3 性能测试', level=2)
    doc.add_paragraph(
        '针对系统核心接口进行性能测试，测试结果如下：'
    )
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    hdr[0].text = '接口'
    hdr[1].text = '数据量'
    hdr[2].text = '响应时间'
    perfs = [
        ('/api/map-observations/', '10000 条（视口过滤）', '< 2 秒'),
        ('/api/species/', '431 个物种', '< 1 秒'),
        ('/api/articles/', '435 篇文章', '< 2 秒'),
        ('/bird/recognize/', '单张图片识别', '< 5 秒'),
    ]
    for api, data, time in perfs:
        row = table.add_row().cells
        row[0].text = api
        row[1].text = data
        row[2].text = time
    doc.add_page_break()


def add_chapter7():
    """第七章 部署与运维"""
    doc.add_heading('七、部署与运维', level=1)

    doc.add_heading('7.1 部署架构', level=2)
    doc.add_paragraph('系统部署于阿里云 ECS 服务器，采用 Nginx + Gunicorn + Django 的标准生产架构：')
    items = [
        'Nginx：处理 HTTPS 终端、静态文件代理和反向代理。',
        'Gunicorn：5 个 Worker 进程 + 2 线程/Worker，通过 Unix Socket 与 Nginx 通信。',
        'Django：业务应用，DEBUG=False，配置数据库连接池（CONN_MAX_AGE=60）。',
        'PostgreSQL：本地数据库服务，PostGIS 扩展已启用。',
        'Systemd：管理 Gunicorn 服务的启停和自动重启。',
    ]
    for item in items:
        doc.add_paragraph(item, style='List Bullet')

    doc.add_heading('7.2 性能优化措施', level=2)
    items = [
        '关闭 DEBUG 模式，避免 SQL 查询日志内存泄漏。',
        '配置 Gunicorn max-requests=500，Worker 定期重启释放内存。',
        '建立 MapObservationCache 缓存表，预计算地图展示数据。',
        '配置 Django 缓存框架，API 响应缓存 5 分钟。',
        '数据库连接池复用（CONN_MAX_AGE=60），减少连接开销。',
        '物种图片下载到本地服务器，避免依赖国外 CDN。',
    ]
    for item in items:
        doc.add_paragraph(item, style='List Bullet')

    doc.add_heading('7.3 访问地址', level=2)
    doc.add_paragraph('公网访问地址：https://3.1415926.love/')
    doc.add_paragraph('后台管理地址：https://3.1415926.love/admin/')
    doc.add_page_break()


def add_chapter8():
    """第八章 创新点与应用价值"""
    doc.add_heading('八、创新点与应用价值', level=1)

    doc.add_heading('8.1 创新点', level=2)

    doc.add_heading('8.1.1 专业监测与公众协同融合', level=3)
    doc.add_paragraph(
        '项目将巡护员、观鸟爱好者和普通公众纳入数据采集链路，公众上报经后台审核后进入平台数据池，'
        '形成"专业监测 + 社会协同"的生态治理闭环，突破了传统监测系统仅服务专业人员的局限。'
    )

    doc.add_heading('8.1.2 空间可视化与生态分析深度结合', level=3)
    doc.add_paragraph(
        '首页将湿地区域、监测点位、监测样线、观鸟记录统一叠加到地图中，'
        '热力图功能支持用户圈选任意区域进行鸟类数量统计，直观呈现鸟类活动热点，辅助巡护决策。'
    )

    doc.add_heading('8.1.3 AI 图像识别嵌入业务流程', level=3)
    doc.add_paragraph(
        '水鸟识别模块不是孤立的模型演示，而是服务于观测上报和物种辅助判断。'
        '系统支持检测阈值和分类阈值设置，使识别过程透明可控，降低公众参与物种辨识的门槛。'
    )

    doc.add_heading('8.1.4 物种数据治理与知识图谱构建', level=3)
    doc.add_paragraph(
        '针对物种别名、同学名重复、图片缺失等真实数据问题，系统实现了图库同步、'
        '图片多源补全（Wikidata/GBIF/iNaturalist）、同学名自动去重和搜索别名兼容等数据治理能力。'
    )

    doc.add_heading('8.1.5 游戏化科普传播', level=3)
    doc.add_paragraph(
        '提供 6 个湿地主题互动游戏，将生态保护知识从静态文字转化为可体验的交互过程，'
        '适合学校、科普馆、湿地公园等场景使用，有效提升公众参与意愿。'
    )

    doc.add_heading('8.2 应用价值', level=2)
    items = [
        '服务黄河流域生态保护和高质量发展国家战略。',
        '为湿地公园、自然保护区提供数字化监测工具。',
        '为学校和科普机构提供生态教育互动平台。',
        '为观鸟社团和公众提供低门槛的参与渠道。',
        '积累的观测数据可为生态研究和政策制定提供数据支撑。',
    ]
    for item in items:
        doc.add_paragraph(item, style='List Bullet')
    doc.add_page_break()


def add_chapter9():
    """第九章 总结与展望"""
    doc.add_heading('九、总结与展望', level=1)

    doc.add_heading('9.1 项目总结', level=2)
    doc.add_paragraph(
        '"黄河生态方舟"围绕黄河湿地鸟类保护场景，设计并实现了一套集观测上报、后台审核、'
        '空间可视化、AI 识别、物种百科、图库科普和互动传播于一体的 Web 平台。'
        '系统具有完整的数据链路、真实的业务逻辑和可运行的线上部署，'
        '体现了软件工程从需求分析到系统实现的完整过程。'
    )
    doc.add_paragraph(
        '项目的核心价值在于：将专业生态监测与公众协同参与连接起来，'
        '让每一次观鸟记录都能成为可管理、可分析、可传播的生态保护力量。'
    )

    doc.add_heading('9.2 不足与改进方向', level=2)
    items = [
        '当前数据量以样例与阶段性观测为主，后续将接入更多实时数据源。',
        'AI 模型训练集有限，需要继续扩展本地鸟类样本以提升识别准确率。',
        '移动端目前偏轻量，后续可完善离线上报、巡护任务分配等功能。',
        '可进一步加入种群趋势预测、异常预警和自动生态报告生成。',
    ]
    for item in items:
        doc.add_paragraph(item, style='List Bullet')

    doc.add_heading('9.3 未来展望', level=2)
    doc.add_paragraph(
        '未来，"黄河生态方舟"将从以下方向持续演进：'
    )
    items = [
        '数据扩展：接入更多湿地区域和鸟类数据，覆盖黄河全流域。',
        '模型优化：持续积累本地样本，提升 AI 识别的物种覆盖度和准确率。',
        '功能深化：加入巡护任务管理、自动报告生成和多端协同能力。',
        '生态价值：为黄河流域生物多样性保护提供长期、可持续的数据支撑。',
    ]
    for item in items:
        doc.add_paragraph(item, style='List Bullet')

    # 结尾
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('— 全文完 —')
    run.font.size = Pt(14)
    run.bold = True


add_chapter3()
add_chapter4()
add_chapter5()
add_chapter6()
add_chapter7()
add_chapter8()
add_chapter9()

# 保存最终文档
output_path = '黄河生态方舟_设计与开发文档.docx'
doc.save(output_path)
print(f'文档已生成: {output_path}')
