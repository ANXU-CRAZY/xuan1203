# -*- coding: utf-8 -*-
"""
严格按照官方模板《2026易智瑞杯 中国大学生GIS软件开发竞赛 项目计划书
（C-GIS应用开发组适用）》生成 docx。

原则：
1. 模板自带的所有固定文字（说明、报名说明、参赛须知、本组报名流程、
   C-GIS应用开发组 基本说明、项目计划书应包括如下内容、章节标题、
   表格的左列字段、页脚联系方式）一字不改；
2. 仅在模板留白处填入"黄河生态方舟"项目内容；
3. 表格结构、字段名、顺序与原模板保持一致。
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


# ---------- 字体与样式 ----------
def set_font(run, name="宋体", size=10.5, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_default_style(doc):
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = Pt(10.5)


def add_para(doc, text, size=10.5, bold=False, font="宋体",
             align=None, indent=False, color=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    if text:
        r = p.add_run(text)
        set_font(r, name=font, size=size, bold=bold, color=color)
    return p


def fill_cell(cell, text, bold=False, font="宋体", size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    if text:
        r = p.add_run(text)
        set_font(r, name=font, size=size, bold=bold)


# ---------- 主体 ----------
def main():
    doc = Document()
    set_default_style(doc)

    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.18)
        section.right_margin = Cm(3.18)

    # ===== 顶部说明（原文不动） =====
    add_para(
        doc,
        "说明：此文件为报名时必须要提交的文件，作为报名的一个重要组成部分不可缺少，"
        "如参赛小组不提交该文档，则报名无效",
    )
    add_para(doc, "")

    # ===== 标题块（居中加粗） =====
    add_para(doc, "2026易智瑞杯", size=18, bold=True, font="黑体",
             align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "中国大学生GIS软件开发竞赛", size=18, bold=True, font="黑体",
             align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "项目计划书", size=18, bold=True, font="黑体",
             align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "")
    add_para(doc, "（C-GIS应用开发组适用）", size=14, bold=False, font="黑体",
             align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "")

    # ===== 团队信息表（原模板就是 5 行 2 列） =====
    info_rows = [
        ("作品名称",
         "黄河生态方舟：基于GeoScene平台的黄河湿地鸟类智能监测与公众协同Web平台"),
        ("报名单位", "（请填写参赛院校全称）"),
        ("团队成员", "（请按报名系统中显示的姓名顺序填写，2~9名）"),
        ("指导老师", "（请填写指导教师姓名，不多于2名）"),
        ("队长及电话", "（请填写队长姓名及联系电话）"),
    ]
    table = doc.add_table(rows=len(info_rows), cols=2)
    table.style = "Table Grid"
    for i, (k, v) in enumerate(info_rows):
        table.rows[i].cells[0].width = Cm(4.5)
        table.rows[i].cells[1].width = Cm(11.5)
        fill_cell(table.rows[i].cells[0], k, bold=True)
        fill_cell(table.rows[i].cells[1], v)
    add_para(doc, "")

    # ===== 报名说明（原文不动） =====
    add_para(doc, "报名说明：")
    add_para(
        doc,
        "1) 以上信息为团队信息，队长的详细联系方式，包括邮箱、快递地址及邮编，"
        "请在报名系统中详尽填写，否则可能没法收取软件及软件许可，以及相关通知；",
    )
    add_para(
        doc,
        "2) 2026年4月15日集中报名截至后，选手仍然可以报名参赛，"
        "但组委会将不再提供参赛软件。",
    )
    add_para(doc, "")

    # ===== 竞赛相关信息（原文不动） =====
    add_para(doc, "竞赛相关信息请见：")
    add_para(doc, "GIS开发竞赛官方网站：http://contest.geoscene.cn/index.html")
    add_para(doc, "技术咨询区：http://zhihu.geoscene.cn/")
    add_para(doc, "GIS开发竞赛官网微博：http://weibo.com/esricontest")
    add_para(doc, "")

    # ===== 参赛须知（原文不动） =====
    add_para(doc, "参赛须知：")
    add_para(
        doc,
        "参赛作品必须是原创作品，并且参赛者均须保证其提交的作品是由其本人或所属参赛团队原创并拥有、"
        "以前从未被发表或发布或许可给第三方发表或发布、以及不损害任何第三方的名誉权、隐私权等任何权利。"
        "参赛作品的原创版权归参赛团队所有，竞赛组委会仅拥有对获奖作品进行展示及推广的权利。"
        "如果提交作品，则意味着接受并遵守参赛要求和参赛规则。",
    )
    add_para(doc, "")
    add_para(doc, "报名截止时间：2026年4月15日")
    add_para(doc, "")

    # ===== 本组报名流程（原文不动） =====
    add_para(doc, "本组报名流程:")
    add_para(
        doc,
        "（1） 在报名系统中注册用户contest.geoscene.cn。每位小组成员均需注册，"
        "为了保证团队队长能够正确填加小组成员，需完整信息（2021年已经在报名系统内注册过的老师同学无需再次注册）。"
        "成员联系方式仅供组委会发送软件申请书、许可及重要紧急情况下联系用，"
        "因此请保证电话号码真实，定期收取邮件，快递信息准确；",
    )
    add_para(
        doc,
        "（2） 选择需要参赛的分组-GIS应用开发组，并依次填加小组成员及指导老师（此项后期可修改）；",
    )
    add_para(
        doc,
        "（3） 将填写好的项目计划书，进行上载（请注意项目计划书文件的大小，尽量不要超过1.5m）；",
    )
    add_para(
        doc,
        "（4） 组委会在收到该文件后，会给予审核，审核通过后，系统通过站内短信通知您的参赛编号，"
        "如审核未通过，您会收到站内短信并获知未通过审核的原因；",
    )
    add_para(
        doc,
        "（5） 以下几种情况可能导致报名审核无法通过：成员在系统中的显示名称与项目计划书不符，"
        "如为网络id等；项目计划书内容缺失；未上载项目计划书；项目计划书计划内容与本组要求不符。",
    )
    add_para(doc, "")

    # ===== 基本说明（原文不动） =====
    add_para(doc, "C-GIS应用开发组 基本说明", bold=True)
    add_para(
        doc,
        "本组采用应用系统开发的方式以利用GIS技术解决实际问题为主线， "
        "综合考察参赛团队和个人的发现问题、分析问题、解决问题的能力；系统设计、开发能力；"
        "新技术探索应用能力；GIS技术综合应用能力；项目管理能力；成果交付能力等。",
    )
    add_para(
        doc,
        "参赛者可使用Web API、Runtime SDK、pro SDK，实现基于ArcGIS/GeoScene平台中的相关产品，"
        "包括Server、Pro、GeoEvent等的二次开发，以及与Portal中可定制的Apps"
        "（Web AppBuilder、Insights、Map Story等）相结合的各种桌面端、移动端、web端应用，"
        "原生移动应用作品可侧重移动用户体验和应用模式创新。"
        "二维、三维以及二三维结合作品皆可。作品服务器端如完全脱离ArcGIS/GeoScene产品平台，将无法参评。"
        "整体作品全部使用开源GIS产品的无法参评。",
    )
    add_para(doc, "运用地图故事模板，讲好一个地图故事，建议参加  A-地图故事组；")
    add_para(doc, "作品不开发，只展示GIS应用设计，建议参加   B-地理设计组")
    add_para(doc, "遥感解译为主的应用作品，建议参加   D-遥感应用组。")
    add_para(doc, "")

    # ===== 项目计划书内容（原模板提示语保留） =====
    add_para(doc, "项目计划书应包括如下内容（请以此为模板填写）：", bold=True)
    add_para(doc, "")

    # ---------- 1 作品概述 ----------
    add_para(doc, "1、 作品概述", bold=True)
    add_para(doc, "作品背景/选题动机/目的")
    add_para(doc, "")
    add_para(doc, "（1）作品背景", bold=True)
    add_para(
        doc,
        "黄河流域是我国生态保护和高质量发展的重大国家战略区，沿黄湿地是东亚-澳大利西亚候鸟迁徙路线上的关键节点，"
        "承载着白枕鹤、大鸨、东方白鹳、黑鹳等珍稀水鸟的繁殖、栖息与越冬功能。"
        "传统的湿地鸟类监测长期面临三个突出问题：一是观鸟数据分散在不同部门、不同保护区与不同志愿组织手中，"
        "缺乏统一的空间数据库；二是公众参与门槛高，巡护员手工填表、爱好者拍照随手发圈，数据无法回流到管理部门；"
        "三是物种知识、识别能力、管理后台与展示前端相互割裂，地图、图库、科普、AI识别没有打通成完整的业务链路。",
        indent=True,
    )
    add_para(doc, "（2）选题动机", bold=True)
    add_para(
        doc,
        "我们希望面向黄河湿地这一具体生态场景，用GIS技术把"
        "“专业监测—公众上报—空间分析—智能识别—科普传播”做成一个统一的Web平台。"
        "GIS不只是底图叠加，而应该成为数据治理、空间分析与决策辅助的核心引擎，"
        "这也是我们选择C-GIS应用开发组的原因。",
        indent=True,
    )
    add_para(doc, "（3）目的", bold=True)
    add_para(
        doc,
        "构建一套基于GeoScene平台的“黄河生态方舟”Web应用系统，"
        "实现湿地区域、监测点位、监测样线、观鸟记录、物种知识、AI识别记录的一体化空间管理与可视化分析，"
        "为湿地保护管理者提供工具，为公众参与生态保护提供入口，"
        "最终输出“可上报—可审核—可分析—可识别—可科普”的一体化生态保护数字底座。",
        indent=True,
    )
    add_para(doc, "")

    # ---------- 2 需求分析 ----------
    add_para(doc, "2、 需求分析", bold=True)
    add_para(doc, "应用领域/实用性分析")
    add_para(doc, "")
    add_para(doc, "（1）应用领域", bold=True)
    for line in [
        "① 自然保护区与湿地公园的鸟类监测与巡护管理；",
        "② 林业、生态环境主管部门的生物多样性数据汇集与展示；",
        "③ 高校、中小学和科普场馆的生态教育与公众科普；",
        "④ 观鸟协会、志愿者组织的众包数据采集与协同治理。",
    ]:
        add_para(doc, line, indent=True)
    add_para(doc, "（2）实用性分析", bold=True)
    items = [
        ("① 管理端实用性。",
         "系统提供基于GeoScene Server的空间数据发布与管理能力，"
         "管理员可在GeoScene Pro中维护湿地边界、监测点位、监测样线等矢量数据，"
         "发布为Feature Service / Map Service，再通过Web端进行编辑、审核与统计，"
         "告别Excel加纸质表格的传统模式。"),
        ("② 巡护端实用性。",
         "巡护员和观鸟爱好者可以通过浏览器或移动浏览器，定位当前位置、选择物种、上传照片，"
         "记录直接进入空间数据库，并按经纬度落到地图上，由管理员在后台审核入库。"),
        ("③ 分析端实用性。",
         "系统支持时间维度筛选、图层叠加、矩形/多边形圈选、热力分析、附近500米观鸟预警、"
         "缓冲区分析等空间分析功能，使观鸟数据真正变成可分析的地理数据。"),
        ("④ 科普端实用性。",
         "物种百科、鸟类图库、科普文章和互动小游戏面向公众用户，"
         "把生态保护从“少数专业人员的工作”扩展为“全民参与”。"),
        ("⑤ AI端实用性。",
         "集成水鸟图像识别模型（YOLO检测加分类双模型），"
         "上传一张鸟类照片即可定位鸟体并给出物种判别，用于辅助新手识鸟和巡护现场速查。"),
        ("⑥ 推广实用性。",
         "整套作品已部署在公网（示例地址http://8.130.88.229），"
         "具备真实可访问性、可演示性和可扩展性，能够直接服务于黄河沿线的湿地公园、保护区与高校实习基地。"),
    ]
    for head, body in items:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)
        r1 = p.add_run(head)
        set_font(r1, bold=True)
        r2 = p.add_run(body)
        set_font(r2)
    add_para(doc, "")

    # ---------- 3 功能设计概述 ----------
    add_para(doc, "3、 功能设计概述", bold=True)
    add_para(doc, "系统架构图/功能模块描述")
    add_para(doc, "")
    add_para(doc, "（1）系统架构图", bold=True)
    add_para(
        doc,
        "系统总体采用B/S架构。GIS空间服务由GeoScene Server发布，"
        "业务逻辑、用户、积分、上报与AI识别由Django + DRF后端处理，"
        "二者通过统一的PostgreSQL/PostGIS空间数据库联通；"
        "前端基于GeoScene Maps SDK for JavaScript加载底图与业务图层，"
        "并叠加业务弹窗、热力分析与统计图表。整体架构层次如下：",
        indent=True,
    )
    arch_lines = [
        "浏览器 / 移动端浏览器 —— 生态大屏、上报、百科、图库、识别页面",
        "        ↓ HTTPS",
        "Nginx 反向代理 + 静态资源托管",
        "        ↓",
        "应用服务层：GeoScene Server（Feature/Map/GP服务） + Gunicorn + Django/DRF（业务API）",
        "        ↓",
        "数据层：PostgreSQL 16 + PostGIS 3 空间数据库（物种、点位、样线、观测、图库、用户）",
        "        ↓",
        "AI与媒体资源层：YOLO检测模型、分类模型与上传影像库",
    ]
    for line in arch_lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.74)
        r = p.add_run(line)
        set_font(r, name="Consolas", size=9)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    add_para(doc, "（2）功能模块描述", bold=True)
    modules = [
        ("M1 生态大屏与地图监测",
         "加载天地图与GeoScene影像底图，叠加湿地区域、监测点位、监测样线、观鸟记录矢量瓦片；"
         "支持图层切换、时间筛选、记录弹窗、保护等级占比图表。"),
        ("M2 空间分析模块",
         "支持矩形/多边形圈选、南北向最小外接矩形热力图、附近500米观鸟预警、观测点缓冲区分析。"),
        ("M3 公众上报模块",
         "提供Web与移动浏览器的鸟类观测上报入口，自动获取经纬度，"
         "支持物种、点位、数量、图片、描述字段，上报后默认进入待审核状态。"),
        ("M4 后台审核模块",
         "管理员对公众上报记录进行审核（pending / approved / rejected），审核通过后进入展示数据池。"),
        ("M5 物种百科与图库",
         "物种中文名、拉丁名、目科、保护等级、IUCN状态、分布习性、封面图、图库图片、相关观测计数与文章数；"
         "支持图片来源回填与同学名去重。"),
        ("M6 科普文章模块",
         "湿地保护、物种科普文章发布与浏览，支持浏览量统计与分类筛选。"),
        ("M7 AI水鸟识别模块",
         "基于YOLO检测加分类双模型对上传图片进行鸟体定位与物种判别，"
         "支持检测阈值与分类阈值参数，识别结果保留到AI识别历史表。"),
        ("M8 用户与积分模块",
         "注册、登录、Token鉴权、积分获取、商品兑换、个人头像、个人资料更新。"),
        ("M9 互动游戏模块",
         "看图识鸟、迁徙、湿地修复、湿地侦探、飞鸟跑酷、浮岛等互动小游戏，用于公众科普。"),
        ("M10 数据导入与维护",
         "通过Django Management命令完成SHP导入、物种元数据同步、Wikimedia图片回填、"
         "保护等级标准化、地图缓存刷新等数据治理任务。"),
    ]
    for name, desc in modules:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)
        r1 = p.add_run(name + "：")
        set_font(r1, bold=True)
        r2 = p.add_run(desc)
        set_font(r2)
    add_para(doc, "")

    # ---------- 4 作品运行环境（严格按原模板 7 行表） ----------
    add_para(doc, "4、 作品运行环境 ", bold=True)
    add_para(doc, "使用软件/操作平台")
    add_para(doc, "")

    env_rows = [
        ("体系结构",
         "B/S（Browser/Server）三层架构：浏览器前端 + 应用服务层 + 空间数据库层。"
         "GIS空间服务通过GeoScene Server发布，业务服务通过Gunicorn + Django/DRF提供。"),
        ("开发平台",
         "Windows 10/11（开发与建库）、Ubuntu 22.04 LTS（生产部署）。"),
        ("开发工具",
         "GeoScene Pro 3.x（数据建库与服务发布）、GeoScene Server 11.x、"
         "GeoScene Maps SDK for JavaScript 4.x、PyCharm Professional、Visual Studio Code、"
         "Git/GitHub、Postman、DBeaver、QGIS（数据校核）、Nginx 1.24、Gunicorn 21、SimpleUI后台美化。"),
        ("开发语言",
         "Python 3.11（后端、AI推理、数据治理脚本）、JavaScript ES2020、HTML5、CSS3（前端）、"
         "SQL/PL-pgSQL（空间查询）、Shell（部署脚本）。"),
        ("运行环境",
         "服务端：Ubuntu 22.04 LTS + Nginx 1.24 + Gunicorn 21 + Python 3.11；"
         "客户端：Chrome、Edge、Firefox等主流浏览器以及iOS / Android移动浏览器；"
         "公网示例地址：http://8.130.88.229。"),
        ("数据库",
         "PostgreSQL 16 + PostGIS 3.4 空间扩展，存储物种信息、湿地区域、监测点位、监测样线、"
         "观测记录、物种图库、用户档案、AI识别历史、积分商城等业务表。"),
        ("其他（可扩充）",
         "AI框架：PyTorch 2.x + Ultralytics YOLO（detector.pt鸟体检测、classifier.pt物种分类）；"
         "主要后端依赖：Django 5.2、Django REST Framework 3.16、djangorestframework-gis 1.2、"
         "django-leaflet、django-cors-headers、django-import-export、Pillow、psycopg2-binary；"
         "空间服务：Feature Service / Map Service / GeoProcessing Service / Mapbox Vector Tile矢量瓦片；"
         "数据治理：Shapefile批量导入、物种元数据同步、Wikimedia Commons图片回填、保护等级标准化。"),
    ]
    env_table = doc.add_table(rows=len(env_rows), cols=2)
    env_table.style = "Table Grid"
    for i, (k, v) in enumerate(env_rows):
        env_table.rows[i].cells[0].width = Cm(3.5)
        env_table.rows[i].cells[1].width = Cm(12.5)
        fill_cell(env_table.rows[i].cells[0], k, bold=True)
        fill_cell(env_table.rows[i].cells[1], v)
    add_para(doc, "")

    # ---------- 5 作品制作周期 ----------
    add_para(doc, "5、 作品制作周期", bold=True)
    add_para(doc, "")
    add_para(
        doc,
        "项目总周期约6个月（2026年1月至2026年6月），划分为8个阶段，"
        "覆盖需求调研、空间数据建库、GIS服务发布、后端开发、前端可视化、AI集成、"
        "部署联调与答辩准备：",
        indent=True,
    )
    schedule = [
        ("第一阶段：需求与数据调研", "2026.01—2026.01",
         "黄河湿地鸟类资料收集、保护区与监测点位调研、用户角色与功能用例梳理、技术选型。"),
        ("第二阶段：空间数据建库", "2026.02—2026.02",
         "在GeoScene Pro中建立物种、湿地区域、监测点位、监测样线要素类，"
         "导入历史观鸟记录与SHP数据，建立PostgreSQL/PostGIS空间数据库。"),
        ("第三阶段：GIS服务发布", "2026.02—2026.03",
         "通过GeoScene Server发布Feature Service、Map Service与GeoProcessing Service，"
         "配置图层符号化与字段别名。"),
        ("第四阶段：后端与API开发", "2026.02—2026.04",
         "Django + DRF业务API、Token鉴权、上报与审核流程、积分体系、"
         "矢量瓦片接口与附近预警接口。"),
        ("第五阶段：前端可视化", "2026.03—2026.05",
         "基于GeoScene Maps SDK for JavaScript实现生态大屏、物种百科、图库、科普文章、"
         "上报页与互动游戏。"),
        ("第六阶段：AI识别集成", "2026.04—2026.05",
         "训练与微调水鸟检测、分类模型，封装识别接口，集成到上报与识别页。"),
        ("第七阶段：部署与联调", "2026.05—2026.05",
         "Nginx + Gunicorn + GeoScene Server部署、性能调优、移动端适配、跨域与安全配置。"),
        ("第八阶段：测试与答辩准备", "2026.06—2026.06",
         "功能回归测试、演示视频录制、PPT制作、作品文档撰写与模拟答辩。"),
    ]
    sch_table = doc.add_table(rows=len(schedule) + 1, cols=3)
    sch_table.style = "Table Grid"
    for j, h in enumerate(["阶段", "时间", "工作内容"]):
        fill_cell(sch_table.rows[0].cells[j], h, bold=True)
    for i, row in enumerate(schedule, start=1):
        for j, val in enumerate(row):
            fill_cell(sch_table.rows[i].cells[j], val, size=10)
    add_para(doc, "")

    # ---------- 6 团队参赛口号或参赛宣言 ----------
    add_para(doc, "6、 团队参赛口号或参赛宣言", bold=True)
    add_para(doc, "")
    add_para(doc, "团队口号：让每一次观鸟，都成为守护黄河的坐标。", indent=True)
    add_para(doc, "")
    add_para(
        doc,
        "参赛宣言：我们用GIS把黄河湿地的鸟类数据、空间分析、AI识别和公众参与连成一张网，"
        "让数据可见，让保护可达，让黄河生态方舟，载着每一只候鸟安全归来。",
        indent=True,
    )
    add_para(doc, "")

    # ---------- 7 学生证（原文不动） ----------
    add_para(doc, "7、 团队成员的学生证信息（图片）", bold=True)
    add_para(
        doc,
        "每位成员的学生证扫描件或清晰翻拍照（请处理成小于100k的jpg文件，或进行拍照）"
        "一起作为图片贴在此处。",
    )
    add_para(doc, "")
    add_para(
        doc,
        "截止作品提交日期，每位小组成员均需保证是在校身份，如大四学生继续升学，"
        "需要提供研究生录取通知书，如暂时无法提供该证明，请在个人学生证信息下说明升学后的单位及开学时间。",
    )
    add_para(doc, "")
    add_para(
        doc,
        "【在此处依次粘贴团队成员（2~9人）的学生证图片，图片下方简要标注姓名、学号、专业、年级。】",
        color=RGBColor(0x99, 0x99, 0x99),
    )
    add_para(doc, "")

    # ---------- 8 团队介绍与分工 ----------
    add_para(doc, "8、 团队介绍（包括个人专业介绍及分工）", bold=True)
    add_para(doc, "")
    add_para(
        doc,
        "团队由来自同一所院校的全日制本科生组成，专业覆盖地理信息科学、计算机科学与技术、"
        "软件工程等方向，具备GIS建库、Web开发、AI模型训练与生态科普内容设计的综合能力。"
        "下方为分工模板，正式提交前请将占位内容替换为实际信息。",
        indent=True,
    )

    members = [
        ("队长 / 项目负责人：xxx",
         [
             "专业年级：xxx大学 xx学院 xxx专业 2023级。",
             "主要分工：项目总体设计与进度管控；GeoScene Pro数据建库与服务发布；"
             "PostGIS空间数据库设计；答辩主讲。",
             "承担模块：M1 生态大屏、M2 空间分析、M10 数据治理。",
         ]),
        ("成员二：xxx",
         [
             "专业年级：xx学院 xxx专业 2023级。",
             "主要分工：Django + DRF后端开发；Token鉴权与权限；上报与审核业务接口；"
             "矢量瓦片与附近预警API。",
             "承担模块：M3 公众上报、M4 后台审核、M8 用户与积分。",
         ]),
        ("成员三：xxx",
         [
             "专业年级：xx学院 xxx专业 2023级。",
             "主要分工：基于GeoScene Maps SDK for JavaScript的前端可视化开发；"
             "图层切换、时间筛选、热力图圈选、统计图表。",
             "承担模块：M1 生态大屏、M2 空间分析（前端部分）、M9 互动游戏。",
         ]),
        ("成员四：xxx",
         [
             "专业年级：xx学院 xxx专业 2023级。",
             "主要分工：AI水鸟识别模型训练与部署；物种百科与图库数据治理；"
             "科普文章与互动游戏内容设计。",
             "承担模块：M5 物种百科与图库、M6 科普文章、M7 AI识别。",
         ]),
    ]
    for title, lines in members:
        add_para(doc, title, bold=True)
        for line in lines:
            add_para(doc, line, indent=True)
        add_para(doc, "")

    add_para(
        doc,
        "（如团队为5人及以上，按以上格式继续补充；"
        "如团队为2~3人，可将上述分工合并到队长与剩余成员名下。）",
        indent=True,
    )
    add_para(doc, "")

    add_para(doc, "指导教师：xxx", bold=True)
    add_para(doc, "所在单位：xx大学 xx学院。", indent=True)
    add_para(
        doc,
        "主要指导内容：项目选题方向、GIS技术路线、系统架构与答辩策略指导。",
        indent=True,
    )
    add_para(doc, "")
    add_para(doc, "")

    # ===== 页脚（原模板格式） =====
    add_para(doc, "                                                                                ")
    p = doc.add_paragraph()
    r = p.add_run("官方网站:  http://contest.geoscene.cn")
    set_font(r, size=10.5)
    p2 = doc.add_paragraph()
    r2 = p2.add_run("\tEmail:  contest@geoscene.cn")
    set_font(r2, size=10.5)

    # 写入到新文件名，避免覆盖正在 Word 中打开的旧版本
    out = r"d:\xuan1203\项目计划书-黄河生态方舟-易智瑞杯C组_V2.docx"
    doc.save(out)
    print(f"已生成：{out}")


if __name__ == "__main__":
    main()
